#!/usr/bin/env python3
"""
Short, plain-language slide deck of the cable-length study (pptx).

Reuses the caches and helpers of cable_length_study.py (run that first) and
draws a few simple figures made for slides, then writes
figs/slides/cable_length_slides.pptx.
"""
import glob
import os

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import uproot
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

import cable_length_study as S

OUT = os.path.join(S.FIGS, "slides")
BLUE, ORANGE, INK, INK2, MUTED = "#2a78d6", "#eb6834", S.INK, S.INK2, S.MUTED
DARK = S.LEN_COLOR[3.0]
EDGES = (1, 31, 62, 63)      # ch 0 is dead in the 1.5 m pulser runs
plt.rcParams.update({"font.size": 13, "axes.titlesize": 14, "axes.labelsize": 13,
                     "legend.fontsize": 11, "xtick.labelsize": 12, "ytick.labelsize": 12})


def _save(fig, name):
    path = os.path.join(OUT, name)
    fig.tight_layout()
    fig.savefig(path, dpi=200)
    plt.close(fig)
    return path


def _fig(w=9, h=5):
    fig, ax = plt.subplots(figsize=(w, h))
    ax.grid(True, color=S.GRID, lw=0.8)
    ax.set_axisbelow(True)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    return fig, ax


# ---------------------------------------------------------------- figures
def fig_spectrum(tab, g):
    """ADC spectrum of one edge channel, short vs long cable (external pulser, 3.3 V)."""
    ch = 1
    fig, ax = _fig()
    for L, col, lw in ((1.0, S.LEN_COLOR[1.0], 2.0), (3.0, DARK, 2.0)):
        t = tab[(tab.cable == "Samtec") & (tab.length == L) & (tab.volt == "3V3") & tab.ok
                & (tab.residue == ch % 3)]
        run = t.run.iloc[0]
        f = [x for x in glob.glob(os.path.join(S.PULSER_DIR, run, "*.root")) if "validation" not in x][0]
        a = uproot.open(f)["hits"].arrays(["hits/vmm", "hits/ch", "hits/adc"], library="np")
        x = a["hits/adc"][(a["hits/vmm"] == S.PULSER_VMM) & (a["hits/ch"] == ch)]
        h = np.bincount(np.clip(x.astype(int), 0, 1023) // 8, minlength=128)
        ax.step(np.arange(128) * 8 + 4, h, where="mid", color=col, lw=lw, label=f"{L:.1f} m cable")
        if L == 3.0:
            med = np.median(x)
            mu = g[(g.cable == "Samtec") & (g.length == L) & (g.volt == "3V3") & (g.ch == ch)].mu.iloc[0]
    ax.set_yscale("log")
    ax.set_xlim(0, 600)
    ax.set_ylim(1, 2e5)
    ax.axvline(med, color=ORANGE, lw=2, ls="--")
    ax.axvline(mu, color=BLUE, lw=2, ls="--")
    ax.text(med + 6, 1.6e5, f"median of all hits = {med:.0f} ADC", color=ORANGE, fontsize=12, va="top")
    ax.text(mu + 6, 4e4, f"Gaussian fit of the\npulse peak = {mu:.0f} ADC", color=BLUE, fontsize=12, va="top")
    ax.annotate("noise hits\n(no pulse)", xy=(115, 300), xytext=(5, 8), fontsize=12, color=INK2,
                arrowprops=dict(arrowstyle="->", color=INK2))
    ax.annotate("real pulses,\n3 m", xy=(262, 500), xytext=(185, 3500), fontsize=12, color=INK2,
                arrowprops=dict(arrowstyle="->", color=INK2))
    ax.annotate("real pulses,\n1 m", xy=(410, 300), xytext=(470, 20), fontsize=12, color=INK2,
                arrowprops=dict(arrowstyle="->", color=INK2))
    ax.set_xlabel("ADC")
    ax.set_ylabel("hits")
    ax.legend(loc="upper right", bbox_to_anchor=(1.0, 0.8))
    ax.set_title(f"Edge channel {ch}, external pulser 3.3 V, Samtec", loc="left", color=INK)
    return _save(fig, "spectrum_edge_channel.png")


def fig_raw_vs_gauss(g):
    """Edge channels / mid-connector vs length: median of all hits vs Gaussian mean."""
    g = g[(g.cable == "Samtec") & (g.volt == "3V3")]
    fig, ax = _fig()
    rows = []
    for L, s in g.groupby("length"):
        s = s.set_index("ch")
        mu = s.mu.where(s.separable)
        rows.append((L, [s.raw_median[c] / np.nanmedian(s.raw_median.reindex(S.MID)) for c in EDGES],
                     [mu[c] / np.nanmedian(mu.reindex(S.MID)) for c in EDGES]))
    L = np.array([r[0] for r in rows])
    raw = np.array([r[1] for r in rows])
    gau = np.array([r[2] for r in rows])
    for i in range(len(EDGES)):
        ax.plot(L, raw[:, i], "-", color=ORANGE, lw=1, alpha=0.35)
        ax.plot(L, gau[:, i], "-", color=BLUE, lw=1, alpha=0.35)
    ax.plot(L, raw.mean(1), "-o", color=ORANGE, lw=3, ms=9, label="median of all hits (simple analysis)")
    ax.plot(L, gau.mean(1), "-o", color=BLUE, lw=3, ms=9, label="Gaussian fit of the pulse peak")
    ax.set_ylim(0, 1.2)
    ax.set_xlabel("total Samtec cable length [m]")
    ax.set_ylabel("edge-channel signal /\nmiddle-of-connector signal")
    ax.legend(loc="lower left")
    ax.set_title("Edge channels 1, 31, 62, 63 (thin: each channel, thick: average)", loc="left", color=INK)
    return _save(fig, "raw_vs_gauss_vs_length.png")


def fig_selftrigger(sn):
    fig, ax = _fig()
    rows = sn["pulser"]["3V3"]["rows"]
    L = np.array(sorted(rows))
    y = np.array([rows[x]["noise_max"] for x in L])
    seen = y > 0
    ax.plot(L[seen], y[seen], "-o", color=BLUE, lw=3, ms=10, label="external pulser")
    ax.plot(L[~seen], np.full((~seen).sum(), 0.1), "v", color=BLUE, ms=11, mfc="white",
            label="external pulser: none seen")
    f = sn["fe55"][1]["noise_max"]
    Lf = sorted(f)
    ax.plot(Lf, [f[x] for x in Lf], "-s", color=INK2, lw=3, ms=10, mfc="white", label="detector (no source)")
    ax.set_yscale("log")
    ax.set_ylim(0.05, 2e4)
    ax.set_xlim(0.9, 3.1)
    ax.set_xlabel("total Samtec cable length [m]")
    ax.set_ylabel("noise hits per second\non the worst edge channel")
    ax.legend(loc="upper left")
    ax.set_title("Hits that are not pulses, edge channels", loc="left", color=INK)
    return _save(fig, "selftrigger_vs_length.png")


def fig_noise(data):
    nb = S.pulser_neighbour_hist()
    edge = list(S.EDGE_ZONE)
    fig, ax = _fig()
    Lp = [L for L in (1.0, 1.5, 2.0, 2.5, 3.0) if ("Samtec", L, "3V3") in nb]
    w = [S.hist_level_width(nb[("Samtec", L, "3V3")])[1] for L in Lp]
    ax.plot(Lp, [np.nanmedian(x[S.MID]) for x in w], "-o", color=BLUE, lw=3, ms=10,
            label="external pulser, middle channels")
    ax.plot(Lp, [np.nanmax(x[edge]) for x in w], "--o", color=BLUE, lw=2, ms=10, mfc="white",
            label="external pulser, worst edge channel")
    Ld, md, ed = [], [], []
    for L in (1.5, 2.0, 2.5):
        s = S._run(data, "fe55", "Samtec", L, 3)
        q = S.hist_level_width(data[s]["h_nb"])[1]
        Ld.append(L)
        md.append(np.nanmedian(q[:, S.MID]))
        ed.append(np.nanmax(np.nanmedian(q, axis=0)[edge]))
    ax.plot(Ld, md, "-s", color=INK2, lw=3, ms=10, label="detector, middle channels")
    ax.plot(Ld, ed, "--s", color=INK2, lw=2, ms=10, mfc="white", label="detector, worst edge channel")
    ax.set_ylim(0, 50)
    ax.set_xlim(0.9, 3.1)
    ax.set_xlabel("total Samtec cable length [m]")
    ax.set_ylabel("noise [ADC]")
    ax.legend(loc="upper left")
    ax.set_title("Electronic noise, same gain (3 mV/fC)", loc="left", color=INK)
    return _save(fig, "noise_vs_length.png")


def fig_edge_extra_loss(g, data):
    """Edge channel signal ratio between lengths, divided by the middle-channel ratio."""
    fig, ax = _fig()
    ax.axhspan(0.995, 1.005, color="#f0efec", zorder=0)
    ax.text(3.08, 1.006, "spread of the\nmiddle channels", fontsize=11, color=MUTED, va="bottom", ha="right")
    gs = g[(g.cable == "Samtec") & (g.volt == "3V3")]
    ref = gs[gs.length == 1.5].set_index("ch")
    ref = ref.mu.where(ref.separable)
    first = True
    for L in (2.0, 2.5, 3.0):
        s = gs[gs.length == L].set_index("ch")
        r = s.mu.where(s.separable) / ref
        r = r / np.nanmedian(r.reindex(S.MID))
        ax.plot(np.full(len(EDGES), L - 0.04), [r[c] for c in EDGES], "o", color=BLUE, ms=10,
                label="external pulser (Gaussian fit), ch 1, 31, 62, 63" if first else None)
        first = False
    s15 = S._run(data, "fe55", "Samtec", 1.5, 1)
    first = True
    for L in (2.0, 2.5):
        s = S._run(data, "fe55", "Samtec", L, 1)
        _, _, per_ch, _, _ = S.ratio_stats(data[s]["peaks"], data[s15]["peaks"])
        ax.plot(np.full(len(S.EDGE), L + 0.04), per_ch[list(S.EDGE)], "s", color=INK2, ms=10, mfc="white",
                mew=2, label="detector Fe55, ch 0, 31, 32, 63" if first else None)
        first = False
    ax.axhline(1, color="#c3c2b7", lw=1.2)
    ax.set_ylim(0.95, 1.03)
    ax.set_xlim(1.8, 3.15)
    ax.set_xticks([2.0, 2.5, 3.0])
    ax.set_xlabel("total Samtec cable length [m], compared with 1.5 m")
    ax.set_ylabel("extra loss of the edge channels\n(1 = same loss as the middle)")
    ax.legend(loc="lower left")
    ax.set_title("Do edge channels lose more signal than the others?", loc="left", color=INK)
    return _save(fig, "edge_extra_loss.png")


def fig_signal_vs_length(sn):
    fig, ax = _fig()
    p = sn["pulser"]["3V3"]
    ax.plot(p["L"], p["rel"], "-o", color=BLUE, lw=3, ms=10, label=f"external pulser: {100 * p['slope']:.0f} % per metre")
    for gain, mk in ((1, "s"), (3, "D")):
        f = sn["fe55"][gain]
        ax.plot(f["L"], f["rel"], mk, color="#c3c2b7", ms=8, mfc="white", mew=1.5,
                label="detector Fe55, before the gain correction" if gain == 1 else None)
        ax.errorbar(f["L"], f["rel_cor"], yerr=f["err_cor"], fmt=mk + "--", color=INK2, lw=2,
                    ms=10, mfc="white", mew=2, capsize=4,
                    label=f"detector Fe55, {gain} mV/fC, gain-corrected: {100 * f['slope_cor']:.0f} % per metre")
    ax.axhline(1, color="#c3c2b7", lw=1.2)
    ax.set_ylim(0.75, 1.1)
    ax.set_xlabel("total Samtec cable length [m]")
    ax.set_ylabel("signal / signal with 1.5 m")
    ax.legend(loc="lower left")
    ax.set_title("Signal of a typical channel vs cable length", loc="left", color=INK)
    return _save(fig, "signal_vs_length.png")


def fig_drift(sn):
    dr = sn["drift"]
    fig, ax = _fig()
    runs = dr["runs"]
    y = np.arange(len(runs))[::-1]
    ax.axvspan(100 * (dr["rate"] - dr["err"]), 100 * (dr["rate"] + dr["err"]), color="#e1e0d9")
    ax.axvline(100 * dr["rate"], color=INK, lw=2)
    ax.axvline(0, color="#c3c2b7", lw=1.2)
    ax.errorbar([100 * r[2] for r in runs], y, xerr=[100 * r[3] for r in runs], fmt="o", color=INK2, ms=9,
                capsize=4, lw=1.5)
    ax.set_yticks(y)
    ax.set_yticklabels([f"{r[0].replace(' (', ' ').replace(')', '')}, {r[1]} mV/fC" for r in runs])
    ax.set_xlim(-8, 5)
    ax.set_xlabel("change of the Fe55 peak during the 10-minute run [% per hour]")
    ax.text(100 * dr["rate"] + 0.2, y[-1] - 0.75, f"average {100 * dr['rate']:.1f} ± {100 * dr['err']:.1f} % per hour",
            fontsize=12, color=INK, va="top")
    ax.set_ylim(-1.3, len(runs) - 0.5)
    ax.set_title("Detector gain during each Fe55 run (7 October)", loc="left", color=INK)
    return _save(fig, "gain_drift.png")


# ---------------------------------------------------------------- slides
W, H = Inches(13.333), Inches(7.5)


def _rgb(hexcol):
    return RGBColor.from_string(hexcol.lstrip("#"))


def _text(slide, x, y, w, h, paras, size=20, color=INK, bold_first=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        bullet = isinstance(p, tuple)
        txt = p[1] if bullet else p
        runs = txt.split("**")
        if bullet:
            r = para.add_run()
            r.text = "•  "
            r.font.size = Pt(size)
            r.font.color.rgb = _rgb(INK)
        for j, chunk in enumerate(runs):
            if not chunk:
                continue
            r = para.add_run()
            r.text = chunk
            r.font.size = Pt(size)
            r.font.color.rgb = _rgb(color)
            r.font.bold = (j % 2 == 1) or (bold_first and i == 0)
        para.space_after = Pt(size * 0.6)
    return tb


def _rule(slide, x, y, w):
    ln = slide.shapes.add_shape(1, x, y, w, Pt(1.2))
    ln.fill.solid()
    ln.fill.fore_color.rgb = _rgb(INK)
    ln.line.fill.background()


def _slide(prs, title, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, Inches(0.6), Inches(0.35), Inches(12.2), Inches(1.0), [title], size=30, bold_first=True)
    _rule(s, Inches(0.6), Inches(1.2), Inches(12.1))
    _text(s, Inches(0.6), Inches(7.0), Inches(9), Inches(0.4),
          ["P2 VMM readout: off-detector cable length (Samtec)"], size=11, color=MUTED)
    _text(s, Inches(12.2), Inches(7.0), Inches(0.8), Inches(0.4), [str(n)], size=11, color=MUTED)
    return s


def _img(slide, path, x, y, w=None, h=None):
    return slide.shapes.add_picture(path, x, y, width=w, height=h)


def build(figs, sn):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    n = 0

    # 1 title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.5),
          ["Choosing the Samtec cable length for P2"], size=40, bold_first=True)
    _rule(s, Inches(0.9), Inches(3.35), Inches(11.5))
    _text(s, Inches(0.9), Inches(3.6), Inches(11.5), Inches(1.5),
          ["Signal loss and the \"broken\" edge channels: external-pulser runs (June) vs Fe55 on the detector (October)",
           "Samtec cables, 1.0 to 3.0 m total length"], size=22, color=INK2)
    _text(s, Inches(0.9), Inches(6.3), Inches(11.5), Inches(0.6), ["8 October 2026"], size=16, color=INK2)

    # 2 question
    n += 1
    s = _slide(prs, "The question", n + 1)
    _text(s, Inches(0.6), Inches(1.5), Inches(12), Inches(5), [
        "P2 will use **Samtec** cables, longer than 1.5 m. **What does each extra length cost?**",
        (1, "**External pulser runs** (June): a pulse generator injects charge into the VMM channels."),
        (1, "**Fe55 runs** (October): a radioactive source gives real signals in the P2 detector."),
        "The external-pulser analysis sees the channels at the **ends of each connector** "
        "(ch 0, 31, 32, 63) \"break down\" with long cables.",
        "The Fe55 analysis does not.",
        "**Which one is right, and why do they differ?**",
    ], size=22)

    # 3 answer
    n += 1
    s = _slide(prs, "The answer in three lines", n + 1)
    _text(s, Inches(0.6), Inches(1.6), Inches(12), Inches(5), [
        (1, "The edge channels **do not lose their signal**. The real pulses are still there, at the same size."),
        (1, "With the external pulser and long cables, the edge channels **also fire on noise**. These extra small "
            "hits pull the average down, so the channel looks broken."),
        (1, "On the detector this doesn't happen: its noise is **3 to 5 times lower**, far below the threshold."),
        "",
        "Every channel loses roughly **10 % of signal per extra metre** of Samtec cable, edges included.",
    ], size=24)

    # 4 spectrum
    n += 1
    s = _slide(prs, "What an edge channel records with a long cable", n + 1)
    _img(s, figs["spectrum"], Inches(0.5), Inches(1.4), w=Inches(8.3))
    _text(s, Inches(9.0), Inches(1.7), Inches(4.0), Inches(5), [
        "With 1 m: one peak, the pulses.",
        "With 3 m: the pulse peak is still there, a bit lower (the cable loss), **plus tens of thousands "
        "of noise hits** at low ADC.",
        "The median of all hits lands on the noise and says the channel lost more than half its signal.",
        "A fit to the pulse peak alone finds the real size.",
    ], size=18)

    # 5 raw vs gauss
    n += 1
    s = _slide(prs, "Fit the pulse peak: the edge channels recover", n + 1)
    _img(s, figs["raw_vs_gauss"], Inches(0.5), Inches(1.4), w=Inches(8.3))
    _text(s, Inches(9.0), Inches(1.7), Inches(4.0), Inches(5), [
        "Orange: the simple analysis. Edge channels drop to **~35 %** from 2.5 m.",
        "Blue: Gaussian fit of the pulse peak. **Flat**: the edge channels keep the same size as with 1 m.",
        "No timing information is used, only the ADC spectrum.",
        "Works at 3.3 V. At 1.8 V the pulse is too close to the noise to separate it from 2.5 m.",
    ], size=18)

    # 6 self-trigger
    n += 1
    s = _slide(prs, "The noise hits appear from 2 m, and only with the external pulser", n + 1)
    _img(s, figs["selftrigger"], Inches(0.5), Inches(1.4), w=Inches(8.3))
    _text(s, Inches(9.0), Inches(1.7), Inches(4.0), Inches(5), [
        "External pulser: **none up to 1.5 m**, then 400 / s at 2 m, 1 000 / s at 2.5 m, 3 600 / s at 3 m.",
        "Detector, no source: **~0.3 / s** at every length, edges like the other channels.",
        "Same channels, same kind of cable: the difference is the setup.",
    ], size=18)

    # 7 noise
    n += 1
    s = _slide(prs, "Why: the external-pulser setup is much noisier", n + 1)
    _img(s, figs["noise"], Inches(0.5), Inches(1.4), w=Inches(8.3))
    _text(s, Inches(9.0), Inches(1.7), Inches(4.0), Inches(5), [
        "Noise measured on channels next to a hit (neighbour trigger), same gain.",
        "Detector: **~4 ADC**, flat with length.",
        "External pulser: **11–22 ADC**, and **~45 ADC** on edge channels from 2 m.",
        "The threshold is ~90 ADC above the baseline: ~2× the edge noise there, so noise crosses it.",
    ], size=18)

    # 8 extra loss
    n += 1
    s = _slide(prs, "Do the edge channels lose more signal? Barely", n + 1)
    _img(s, figs["edge_extra"], Inches(0.5), Inches(1.4), w=Inches(8.3))
    _text(s, Inches(9.0), Inches(1.7), Inches(4.0), Inches(5), [
        "Each channel compared with itself at 1.5 m, then with the middle channels.",
        "Fe55: edges within **0.7 %** of the middle channels.",
        "External pulser: edges **0–3 %** lower, worst on ch 63.",
        "Small either way: no edge channel is lost.",
    ], size=18)

    # 9 signal vs length
    n += 1
    s = _slide(prs, "How much signal does a longer cable cost?", n + 1)
    _img(s, figs["signal"], Inches(0.5), Inches(1.4), w=Inches(8.3))
    pf = sn["pulser"]["3V3"]
    f1, f3 = sn["fe55"][1], sn["fe55"][3]
    _text(s, Inches(9.0), Inches(1.7), Inches(4.0), Inches(5), [
        f"External pulser: **{-100 * pf['slope']:.0f} % per metre**.",
        f"Fe55, corrected for the detector gain drift (next slide): **{-100 * f1['slope_cor']:.0f}–"
        f"{-100 * f3['slope_cor']:.0f} % per metre**.",
        f"Without the correction Fe55 gives {-100 * f1['slope']:.0f}–{-100 * f3['slope']:.0f} % per metre: "
        "the lengths were measured one after the other while the gain was falling.",
        "The two methods agree: **~10 % per metre**, the same on all channels.",
    ], size=18)

    # 10 gain drift
    n += 1
    dr = sn["drift"]
    s = _slide(prs, "How the detector gain drift was measured", n + 1)
    _img(s, figs["drift"], Inches(0.5), Inches(1.4), w=Inches(8.3))
    _text(s, Inches(9.0), Inches(1.7), Inches(4.0), Inches(5), [
        "Each Fe55 run lasts 10 minutes and is written as ten 1-minute files.",
        "The Fe55 peak is measured in each file: inside **8 of 9 runs it goes down**.",
        f"Together: **{100 * dr['rate']:.1f} ± {100 * dr['err']:.1f} % per hour**, the same in every run.",
        "Check: the same Hitachi cable measured twice, 69 min apart, lost 3.6 % (−3.1 % per hour).",
        "The correction assumes the drift continued at the same rate between runs.",
    ], size=17)

    # 11 decision table
    n += 1
    s = _slide(prs, "Samtec lengths above 1.5 m: what each one costs", n + 1)
    lens = (2.0, 2.5, 3.0)
    pr = dict(zip(pf["L"], pf["rel"]))
    c1, c3 = dict(zip(f1["L"], f1["rel_cor"])), dict(zip(f3["L"], f3["rel_cor"]))
    pct = lambda v: f"−{100 * (1 - v):.0f} %"
    rows = [["Total Samtec length", "2.0 m", "2.5 m", "3.0 m"],
            ["Signal lost vs 1.5 m, external pulser"] + [pct(pr[L]) for L in lens],
            ["Signal lost vs 1.5 m, Fe55 (gain-corrected)"]
            + [f"{pct(c1[L])} / {pct(c3[L])}" if L in c1 else "not measured" for L in lens],
            ["Edge channels on the detector", "same as the others", "same as the others", "not measured"],
            ["Noise on the detector", "unchanged (~4 ADC)", "unchanged (~4 ADC)", "not measured"]]
    tbl = s.shapes.add_table(len(rows), 4, Inches(0.6), Inches(1.6), Inches(12.1), Inches(3.4)).table
    tbl.columns[0].width = Inches(4.6)
    for j in range(1, 4):
        tbl.columns[j].width = Inches(2.5)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = _rgb("#ffffff" if i else "#e9e9e9")
            c.text = val
            for para in c.text_frame.paragraphs:
                for r in para.runs:
                    r.font.size = Pt(17)
                    r.font.bold = i == 0 or j == 0
                    r.font.color.rgb = _rgb(INK)
    _text(s, Inches(0.6), Inches(5.3), Inches(12.1), Inches(1.6), [
        "Fe55 values: 1 mV/fC / 3 mV/fC. Roughly **5 % of signal per extra 0.5 m**; no channel is lost at the "
        "connector edges up to 2.5 m on the detector.",
    ], size=17)

    # 12 caveats
    n += 1
    s = _slide(prs, "What is not settled", n + 1)
    _text(s, Inches(0.6), Inches(1.5), Inches(12), Inches(5), [
        (1, "On the detector, the cables were measured **up to 2.5 m**. 3.0 m was only measured with the "
            "external pulser, and the cable type of that run is not written in the logbook (assumed Samtec)."),
        (1, "**Where the external-pulser noise comes from** is not known (pulser connection, extra 50 cm cable and "
            "joint, grounding of the pulser setup). It does not appear on the detector."),
        (1, "The Fe55 gain correction assumes a constant drift between runs. Measuring the lengths in mixed order "
            "would remove that assumption."),
    ], size=21)

    # 13 conclusions
    n += 1
    s = _slide(prs, "For P2 with Samtec cables", n + 1)
    _text(s, Inches(0.6), Inches(1.6), Inches(12), Inches(5), [
        (1, "**No channel is lost** at the connector edges because of the cable length."),
        (1, "Each extra metre costs **~10 % of the signal**, on every channel: keep the cable as short as the "
            "integration allows."),
        (1, "On the detector, the noise does not change between 1.5 and 2.5 m."),
    ], size=24)

    # 14 backup: optional checks
    n += 1
    s = _slide(prs, "Backup: optional checks", n + 1)
    _text(s, Inches(0.6), Inches(1.5), Inches(12), Inches(5), [
        "Not needed for the length choice; they would explain the external-pulser noise.",
        (1, "External pulser, long cable, neighbour trigger off: which channel really oscillates?"),
        (1, "Pulser disconnected, cables left on; add a ground strap; remove the 50 cm joint."),
        (1, "Fe55 at 3.0 m on the detector, and the lengths in mixed order (1.5 → 2.5 → 2.0 → 1.5 m)."),
    ], size=20)

    path = os.path.join(OUT, "cable_length_slides.pptx")
    prs.save(path)
    return path


def main():
    os.makedirs(OUT, exist_ok=True)
    tab = S.pulser_table()
    data = S.fe55_cache()
    g = S.pulser_peak_table(tab)
    sn = S.samtec_numbers(tab, data)
    figs = dict(spectrum=fig_spectrum(tab, g), raw_vs_gauss=fig_raw_vs_gauss(g), selftrigger=fig_selftrigger(sn),
                noise=fig_noise(data), edge_extra=fig_edge_extra_loss(g, data), signal=fig_signal_vs_length(sn), drift=fig_drift(sn))
    print("wrote", build(figs, sn))


if __name__ == "__main__":
    main()
